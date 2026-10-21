"""Export every Kreyòl string in Paskle to one spreadsheet for editing.
Usage (from the repo root): python3 tools/kreyol_export.py [out.xlsx]
Writes kreyol-strings.xlsx by default. Edit the "Your edit" column, then run tools/kreyol_import.py on the file."""
import json, os, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# Files that hold Kreyòl text by hand. Generated files (data/cna-a.js, data/ma-cna.js, data/permit-a.js) are rebuilt from the content/ sources.
SOURCES = ["index.html", "data/hha-a.js", "data/hha-b.js", "data/hha-c.js", "data/hha-d.js",
           "content/cna/cna_src.py", "content/permit/permit_src.py", "content/permit/permit_faq.py"]

def quoted(s):
    return '"' + json.dumps(s, ensure_ascii=False)[1:-1] + '"'

def area(where):
    w = where[0]
    for key, name in (("App screens", "App screens"), ("Glossary", "Glossary"), ("HHA FAQ", "HHA FAQ"), ("Permit FAQ", "Permit FAQ"),
                      ("PÈMI", "Permit"), ("CNA", "CNA"), ("HHA", "HHA"), ("CDL", "CDL")):
        if w.startswith(key): return name
    return "Other"

def main(out):
    items = json.loads(subprocess.check_output(["node", os.path.join(ROOT, "tools", "kreyol-extract.js"), ROOT]))
    texts = {f: open(os.path.join(ROOT, f), encoding="utf-8").read() for f in SOURCES}
    rows, skipped = [], 0
    for it in items:
        a = area(it["where"])
        if a == "CDL":  # CDL track is not live yet
            continue
        files = [f for f in SOURCES if quoted(it["ht"]) in texts[f]]
        if not files:
            skipped += 1  # assembled by the app from pieces (e.g. chapter intros); not editable as one string
            continue
        rows.append((a, "; ".join(it["where"]), it["en"], it["ht"], ", ".join(files)))
    order = ["App screens", "Permit", "Permit FAQ", "CNA", "HHA", "HHA FAQ", "Glossary", "Other"]
    rows.sort(key=lambda r: order.index(r[0]))

    wb = Workbook()
    ws = wb.active; ws.title = "Kreyòl"
    head = ["ID", "Area", "Where it appears", "English", "Kreyòl now", "Your edit", "Comment", "File"]
    ws.append(head)
    for i, r in enumerate(rows, 1):
        ws.append([f"K{i:04d}", r[0], r[1], r[2], r[3], None, None, r[4]])
    hfont, base = Font(name="Arial", bold=True, color="FFFFFF"), Font(name="Arial", size=10)
    for c in ws[1]:
        c.font = hfont; c.fill = PatternFill("solid", fgColor="13213D"); c.alignment = Alignment(vertical="center", wrap_text=True)
    edit_fill = PatternFill("solid", fgColor="FFF4CC")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = base; c.alignment = Alignment(vertical="top", wrap_text=True)
        row[5].fill = edit_fill; row[6].fill = edit_fill
        row[0].font = Font(name="Arial", size=10, color="666666"); row[7].font = Font(name="Arial", size=9, color="666666")
        row[2].font = Font(name="Arial", size=9, color="444444")
    for col, w in zip("ABCDEFGH", (8, 13, 34, 48, 48, 48, 24, 22)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "E2"; ws.auto_filter.ref = f"A1:H{ws.max_row}"

    rd = wb.create_sheet("Read me", 0)
    lines = [
        ("Paskle: Kreyòl strings for review", True),
        ("", False),
        ("How to edit", True),
        ("1. Go to the Kreyòl tab. Each row is one Kreyòl text, with its English beside it.", False),
        ("2. To change a text, type the full corrected Kreyòl in the yellow \"Your edit\" column. Leave it empty to keep the text as it is.", False),
        ("3. Use the \"Comment\" column for questions or notes. Comments are not put into the app.", False),
        ("4. Keep anything in curly brackets exactly as it is, for example {n} or {p}. The app fills those with numbers.", False),
        ("5. Keep the meaning the same as the English. The facts come from official sources and are cited in the app.", False),
        ("6. Send the file back to Claude, or run: python3 tools/kreyol_import.py kreyol-strings.xlsx", False),
        ("", False),
        ("Example", True),
        ("Kreyòl now:  Lave men w anvan ou manyen kliyan an.", False),
        ("Your edit:   Lave men ou anvan ou manyen kliyan an.", False),
        ("", False),
        ("Notes", True),
        ("• A text that appears in several places is listed once. Editing it changes it everywhere it appears.", False),
        ("• Filter the Area column to work one part at a time (App screens, Permit, CNA, HHA, FAQ, Glossary).", False),
        ("• Official terms for the permit come from the RMV's Kreyòl manual (April 2026): pèmi apranti, Egzamen Pratik, kafou won, dwa de pasaj.", False),
        (f"• {len(rows)} texts listed. The CDL track is not live yet and is left out.", False),
    ]
    for t, bold in lines:
        rd.append([t]); rd.cell(rd.max_row, 1).font = Font(name="Arial", size=14 if t.startswith("Paskle") else 11, bold=bold)
    rd.column_dimensions["A"].width = 120
    wb.save(out)
    print(f"wrote {out}: {len(rows)} rows, {skipped} assembled strings left out")

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "kreyol-strings.xlsx")
