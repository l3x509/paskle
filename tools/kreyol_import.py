"""Put Kreyòl edits from the review spreadsheet back into the source files, then rebuild.
Usage (from the repo root): python3 tools/kreyol_import.py kreyol-strings.xlsx [--dry-run]
Reads the "Kreyòl" tab: every row with text in "Your edit" that differs from "Kreyòl now" replaces that exact
Kreyòl text in the files that hold it. Then it rebuilds the CNA and permit data and runs the checks."""
import json, os, re, subprocess, sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCES = ["index.html", "data/hha-a.js", "data/hha-b.js", "data/hha-c.js", "data/hha-d.js",
           "content/cna/cna_src.py", "content/permit/permit_src.py", "content/permit/permit_faq.py"]
quoted = lambda s: '"' + json.dumps(s, ensure_ascii=False)[1:-1] + '"'
holes = lambda s: sorted(re.findall(r"\{\w+\}", s))

def main(path, dry):
    ws = load_workbook(path)["Kreyòl"]
    head = [c.value for c in ws[1]]
    ix = {h: head.index(h) for h in ("ID", "Kreyòl now", "Your edit")}
    edits, problems = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        old, new = row[ix["Kreyòl now"]], row[ix["Your edit"]]
        if new is None or not str(new).strip(): continue
        new = " ".join(str(new).split())  # tidy spaces and line breaks
        if new == old: continue
        if holes(new) != holes(old):
            problems.append(f"{row[ix['ID']]}: keep the {{…}} parts unchanged ({', '.join(holes(old)) or 'none'})"); continue
        edits.append((row[ix["ID"]], old, new))
    seen = {}
    for rid, old, new in edits:
        if old in seen and seen[old][1] != new:
            problems.append(f"{rid} and {seen[old][0]}: same Kreyòl text, two different edits")
        seen[old] = (rid, new)
    texts = {f: open(os.path.join(ROOT, f), encoding="utf-8").read() for f in SOURCES}
    done = 0
    for rid, old, new in edits:
        hit = False
        for f in SOURCES:
            q = quoted(old)
            if q in texts[f]:
                texts[f] = texts[f].replace(q, quoted(new)); hit = True
        if hit: done += 1
        else: problems.append(f"{rid}: text not found (already changed?): {old[:60]}")
    print(f"{len(edits)} edits read, {done} applied")
    for p in problems: print("  CHECK", p)
    if dry or not done: return
    for f, t in texts.items():
        open(os.path.join(ROOT, f), "w", encoding="utf-8").write(t)
    for cmd in (["python3", "content/cna/build_cna.py"], ["python3", "content/permit/build_permit.py"],
                ["node", "tools/split_web.js"], ["node", "tools/i18n-check.js"]):
        r = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True)
        print("$", " ".join(cmd), "->", "ok" if r.returncode == 0 else "FAILED");
        if r.returncode: print(r.stdout[-1500:], r.stderr[-1500:])
    print("Now bump VERSION in sw.js, test, and push.")

if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    main(args[0] if args else "kreyol-strings.xlsx", "--dry-run" in sys.argv)
